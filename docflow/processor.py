"""
Document processor implementation.
Handles document processing and metadata extraction.
"""

import yaml
import asyncio
from pathlib import Path
from typing import Dict, List, Optional, Any, Union
from .models import BaseModel as BaseAIModel  # Alias for backward compatibility
from .models import ModelRegistry
import logging
import re
from typing import BinaryIO
import docx
from PyPDF2 import PdfReader
import io
from datetime import datetime
import dateutil.parser
from PIL import Image
import pdf2image
import tempfile

logger = logging.getLogger(__name__)


class DocumentProcessor:
    def __init__(
        self,
        rules_file: str = "config/default_rules.yaml",
        ai_model: Optional[str] = None,
    ):
        self.rules = self._load_rules(rules_file)
        self.supported_formats = {".pdf", ".docx", ".txt", ".jpg", ".jpeg", ".png"}

        # Track temporary files generated during processing for cleanup
        self._temp_files: list[Path] = []

        # Get available models
        self.available_models = ModelRegistry.list_models()
        logger.info(f"Available AI models: {list(self.available_models.keys())}")

        # Initialize AI model based on preference
        self.ai_model = self._initialize_ai_model(ai_model)

    def _process_pdf_for_vision(self, file_path: str) -> str:
        """Convert PDF to image for vision model processing"""
        try:
            # Convert first page of PDF to image
            images = pdf2image.convert_from_path(file_path, first_page=1, last_page=1)
            if not images:
                raise ValueError("Failed to convert PDF to image")

            # Save the image to a temporary file
            with tempfile.NamedTemporaryFile(suffix=".jpg", delete=False) as temp_file:
                images[0].save(temp_file.name, "JPEG")
                temp_path = Path(temp_file.name)
                self._temp_files.append(temp_path)
                return temp_path.as_posix()
        except Exception as e:
            logger.error(f"Error converting PDF to image: {e}")
            raise

    def _extract_text_with_ocr(self, file_path: str) -> str:
        """Extract text using OCR when available"""
        try:
            import docling
        except Exception as e:
            logger.warning(f"OCR requested but docling is unavailable: {e}")
            return ""

        try:
            document = docling.Document(str(file_path))
            return document.extract_text(ocr=True)
        except Exception as e:
            logger.warning(f"OCR extraction failed for {file_path}: {e}")
            return ""

    def _process_image(
        self, file_path: str, use_ocr: bool = False
    ) -> tuple[str, Optional[str]]:
        """Process image file and return extracted text (if OCR) and image path"""
        try:
            # For image files, we'll use them directly for vision models
            # and optionally perform OCR
            image_path = file_path
            extracted_text = self._extract_text_with_ocr(file_path) if use_ocr else ""
            return extracted_text, image_path
        except Exception as e:
            logger.error(f"Error processing image: {e}")
            raise

    def _extract_text_from_pdf(
        self, file_path: str, use_ocr: bool = False, convert_to_img: bool = False
    ) -> tuple[str, Optional[str]]:
        """Extract text from PDF file and convert to image if needed"""
        try:
            # Extract text
            with open(file_path, "rb") as file:
                pdf = PdfReader(file)
                text = ""
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"

            if use_ocr:
                ocr_text = self._extract_text_with_ocr(file_path)
                if ocr_text:
                    text = f"{text}\n{ocr_text}" if text else ocr_text

            # Convert to image for vision models when requested
            image_path = (
                self._process_pdf_for_vision(file_path) if convert_to_img else None
            )

            return text, image_path
        except Exception as e:
            logger.error(f"Error extracting text from PDF: {e}")
            raise

    def _extract_text_from_docx(self, file_path: str) -> tuple[str, None]:
        """Extract text from DOCX file"""
        try:
            doc = docx.Document(file_path)
            return "\n".join([paragraph.text for paragraph in doc.paragraphs]), None
        except Exception as e:
            logger.error(f"Error extracting text from DOCX: {e}")
            raise

    def _extract_text_from_txt(self, file_path: str) -> tuple[str, None]:
        """Extract text from TXT file"""
        try:
            with open(file_path, "r", encoding="utf-8") as file:
                return file.read(), None
        except Exception as e:
            logger.error(f"Error extracting text from TXT: {e}")
            raise

    def _extract_text(
        self, file_path: str, use_ocr: bool = False, convert_to_img: bool = False
    ) -> tuple[str, Optional[str]]:
        """Extract text from document based on file type"""
        path = Path(file_path)
        suffix = path.suffix.lower()

        if suffix == ".pdf":
            return self._extract_text_from_pdf(
                file_path, use_ocr=use_ocr, convert_to_img=convert_to_img
            )
        elif suffix == ".docx":
            return self._extract_text_from_docx(file_path)
        elif suffix == ".txt":
            return self._extract_text_from_txt(file_path)
        elif suffix in {".jpg", ".jpeg", ".png"}:
            return self._process_image(file_path, use_ocr=use_ocr)
        else:
            raise ValueError(f"Unsupported file format: {suffix}")

    def _classify_document(self, text: str) -> str:
        """Classify document based on content and comprehensive rules"""
        max_score = 0
        best_type = "unknown"

        if logger.isEnabledFor(logging.DEBUG):
            logger.debug("Starting document classification")

        for doc_type, type_config in self.rules.get("rules", {}).items():
            keywords = type_config.get("keywords", [])

            # Calculate match score based on keyword presence and metadata fields
            score = 0
            matched_keywords = []

            # Check for keywords
            for keyword in keywords:
                if keyword.lower() in text.lower():
                    score += 1
                    matched_keywords.append(keyword)

            # Check for metadata field patterns if keywords match
            if score > 0:
                metadata_fields = type_config.get("metadata_fields", {})
                for field_name, field_config in metadata_fields.items():
                    pattern = field_config.get("pattern", "")
                    if re.search(pattern, text, re.IGNORECASE | re.MULTILINE):
                        score += 0.5

            if logger.isEnabledFor(logging.DEBUG):
                logger.debug(
                    f"Document type '{doc_type}' scored {score} with keywords: {matched_keywords}"
                )

            if score > max_score:
                max_score = score
                best_type = doc_type

        if logger.isEnabledFor(logging.DEBUG):
            logger.debug(f"Final classification: {best_type} with score {max_score}")
        return best_type

    def process_document(
        self, file_path: str, use_ocr: bool = False, convert_to_img: bool = False
    ) -> Dict:
        """Process a document with the selected AI model and extract information."""
        try:
            path = Path(file_path)
            if path.suffix.lower() not in self.supported_formats:
                raise ValueError(f"Unsupported file format: {path.suffix}")

            # Extract text and get image path if applicable
            text, image_path = self._extract_text(
                file_path, use_ocr=use_ocr, convert_to_img=convert_to_img
            )
            if logger.isEnabledFor(logging.DEBUG):
                logger.debug(
                    f"Extracted text length: {len(text)}, image path: {image_path}"
                )

            # Use AI model for enhanced extraction
            result = {
                "file_name": path.name,
                "text_length": len(text),
                "text_content": text,
            }

            try:
                logger.info(
                    f"Processing with AI model: {self.ai_model.__class__.__name__}"
                )
                ai_analysis = self.ai_model.extract_information(
                    text=text, image_path=image_path
                )
                result["ai_analysis"] = ai_analysis

                if ai_analysis.get("success"):
                    # Extract fields from schema-based response
                    raw_analysis = ai_analysis.get("raw_analysis", {})

                    # Add core fields
                    result["document_type"] = raw_analysis.get("doctype")
                    if raw_analysis.get("date"):
                        result["date"] = raw_analysis["date"]
                    if raw_analysis.get("title"):
                        result["title"] = raw_analysis["title"]
                    if raw_analysis.get("reference"):
                        result["reference"] = raw_analysis["reference"]

                    # Add metadata
                    meta = raw_analysis.get("meta", {})
                    if meta:
                        # Add dates
                        if "dates" in meta:
                            result["dates"] = meta["dates"]

                        # Add amounts
                        if "amounts" in meta:
                            for key, value in meta["amounts"].items():
                                if value is not None:  # Only add non-null values
                                    result[f"{key}_amount"] = value

                        # Add entities
                        if "entities" in meta:
                            result["entities"] = meta["entities"]

                else:
                    logger.warning(
                        f"AI analysis failed: {ai_analysis.get('error', 'Unknown error')}"
                    )

            except Exception as e:
                logger.error(f"AI model analysis failed: {e}", exc_info=True)
                result["ai_analysis"] = {
                    "success": False,
                    "error": str(e),
                    "model_name": self.ai_model.__class__.__name__,
                }

            # Perform basic classification if needed
            if "document_type" not in result or not result["document_type"]:
                try:
                    classification = self._classify_document(text)
                    if isinstance(classification, tuple):
                        doc_type, confidence = classification
                    else:
                        doc_type = classification
                        confidence = 1.0
                    result["document_type"] = doc_type
                    result["classification_confidence"] = confidence
                except Exception as e:
                    logger.error(f"Classification failed: {e}")
                    result["document_type"] = "unknown"
                    result["classification_confidence"] = 0.0

            logger.info(f"Successfully processed document: {path.name}")
            return result

        except Exception as e:
            logger.error(f"Error processing document: {e}", exc_info=True)
            return {
                "error": str(e),
                "file_name": path.name if "path" in locals() else None,
                "document_type": "unknown",  # Add default document type for error cases
            }
        finally:
            for temp_path in self._temp_files:
                try:
                    temp_path.unlink(missing_ok=True)
                except Exception as cleanup_error:
                    logger.warning(
                        f"Failed to remove temp file {temp_path}: {cleanup_error}"
                    )
            self._temp_files.clear()

    def _initialize_ai_model(self, model_name: Optional[str]) -> BaseAIModel:
        """Initialize the specified AI model with proper error handling"""
        if logger.isEnabledFor(logging.DEBUG):
            logger.debug(f"Initializing AI model: {model_name}")

        try:
            # Get model class from registry
            model_class = (
                ModelRegistry.get_model(model_name)
                if model_name
                else ModelRegistry.get_model("fallback")
            )

            if not model_class:
                logger.warning(f"Model {model_name} not found, using fallback")
                model_class = ModelRegistry.get_model("fallback")

            if not model_class.is_available():
                raise Exception(f"Model {model_name} is not available")

            return model_class()

        except Exception as e:
            logger.error(f"Error initializing model '{model_name}': {e}", exc_info=True)
            # Always fall back to the fallback model
            return ModelRegistry.get_model("fallback")()

    def get_supported_models(self) -> Dict[str, str]:
        """Return dictionary of available models and their descriptions"""
        return self.available_models

    def _load_rules(self, rules_file: str) -> Dict:
        try:
            with open(rules_file, "r") as f:
                return yaml.safe_load(f)
        except Exception as e:
            logger.error(f"Error loading rules file: {e}")
            return {}

    def _extract_metadata(self, text: str, doc_type: str) -> Dict:
        """Extract metadata using patterns from rules"""
        metadata = {}
        fields = (
            self.rules.get("rules", {}).get(doc_type, {}).get("metadata_fields", {})
        )

        if logger.isEnabledFor(logging.DEBUG):
            logger.debug(f"Attempting to extract metadata for type '{doc_type}'")
            logger.debug(f"Available fields: {fields}")

        for field_name, field_config in fields.items():
            pattern = field_config.get("pattern", "")
            if field_name == "total_amount":
                pattern = r"(?i)(?:total\s+net|amount\s+due|total|amount|sum|betrag|summe|rechnungsbetrag)[\s:]*[$€£]?\s*([\d.,]+(?:[\.,]\d{2})?)\s*[$€£]?"
            elif field_name == "date":
                pattern = r"(?i)(?:invoice\s+date|date|datum|belegdatum)[\s:]*(\d{1,2}[-/\.]\d{1,2}[-/\.]\d{4})"
            elif field_name == "invoice_number":
                # Updated pattern to handle GC-2024/12/01 format
                pattern = r"(?i)(?:Invoice\s+No\.?:|Rechnung\s+Nr\.?:|Rechnungsnummer:?|Invoice\s+number:?)\s*((?:[A-Za-z]{1,4}[-]?\d{4}[-/]\d{1,2}[-/]\d{1,2})|(?:[A-Za-z0-9][-A-Za-z0-9/]*[A-Za-z0-9]))"

            if logger.isEnabledFor(logging.DEBUG):
                logger.debug(f"Using pattern for {field_name}: {pattern}")

            matches = re.findall(pattern, text, re.MULTILINE | re.IGNORECASE)
            if matches:
                value = matches[0]
                if isinstance(value, tuple):
                    value = value[0]
                try:
                    field_type = field_config.get("type", "String")
                    converted_value = self._convert_value(value, field_type)
                    metadata[field_name] = converted_value
                    if logger.isEnabledFor(logging.DEBUG):
                        logger.debug(
                            f"Field {field_name}: found and converted to {field_type}"
                        )
                except Exception as e:
                    logger.error(f"Error converting field {field_name}: {e}")
                    continue

        return metadata

    def _convert_value(self, value: str, field_type: str) -> Any:
        """Convert extracted value to the specified type"""
        if field_type == "Date":
            return self._parse_date(value)
        elif field_type == "Float":
            return self._parse_float(value)
        elif field_type == "Int":
            return int(float(self._parse_float(value)))
        else:  # String or unknown type
            return value

    def _parse_date(self, date_str: str) -> str:
        """Parse date string to YYYY-MM-DD format"""
        try:
            # Clean up the date string
            date_str = date_str.strip()
            # Parse the date string using dateutil
            parsed_date = dateutil.parser.parse(date_str)
            # Validate year is reasonable (between 1900 and 2100)
            if parsed_date.year < 1900 or parsed_date.year > 2100:
                raise ValueError(f"Invalid year in date: {parsed_date.year}")
            return parsed_date.strftime("%Y-%m-%d")
        except Exception as e:
            logger.error(f"Error parsing date {date_str}: {e}")
            raise ValueError(f"Invalid date format: {date_str}")

    def _parse_float(self, amount_str: str) -> float:
        """Parse string amount to float, handling different number formats"""
        try:
            # Remove any currency symbols and whitespace
            cleaned = (
                amount_str.strip().replace("$", "").replace("€", "").replace("£", "")
            )

            # Handle European number format (1.234,56)
            if "." in cleaned and "," in cleaned:
                # Remove thousands separator (dot) and replace decimal comma with dot
                cleaned = cleaned.replace(".", "").replace(",", ".")
            # Handle cases where comma is used as decimal separator
            elif "," in cleaned and "." not in cleaned:
                cleaned = cleaned.replace(",", ".")

            return float(cleaned)
        except Exception as e:
            logger.error(f"Error parsing float {amount_str}: {e}")
            return 0.0

    # =========================================================================
    # Async Methods
    # =========================================================================

    async def process_document_async(
        self, file_path: str, use_ocr: bool = False, convert_to_img: bool = False
    ) -> Dict:
        """
        Async version of process_document.

        Processes a document with the selected AI model and extract information.
        This method runs the synchronous process_document in a thread pool to enable
        concurrent processing without blocking the event loop.

        Args:
            file_path: Path to the document file
            use_ocr: Whether to use OCR for image-based documents
            convert_to_img: Whether to convert PDF to image for vision models

        Returns:
            Dictionary containing extraction results
        """
        loop = asyncio.get_event_loop()
        return await loop.run_in_executor(
            None,  # Uses default executor
            self.process_document,
            file_path,
            use_ocr,
            convert_to_img,
        )

    async def process_batch_async(
        self,
        file_paths: List[str],
        use_ocr: bool = False,
        convert_to_img: bool = False,
        max_concurrency: int = 4,
    ) -> List[Dict]:
        """
        Process multiple documents concurrently with limited concurrency.

        Args:
            file_paths: List of document file paths
            use_ocr: Whether to use OCR for image-based documents
            convert_to_img: Whether to convert PDF to image
            max_concurrency: Maximum number of concurrent processing tasks

        Returns:
            List of extraction result dictionaries
        """
        semaphore = asyncio.Semaphore(max_concurrency)

        async def process_with_limit(file_path: str) -> Dict:
            async with semaphore:
                try:
                    return await self.process_document_async(
                        file_path, use_ocr=use_ocr, convert_to_img=convert_to_img
                    )
                except Exception as e:
                    logger.error(f"Failed to process {file_path}: {e}")
                    return {
                        "error": str(e),
                        "file_name": Path(file_path).name,
                        "document_type": "unknown",
                        "success": False,
                    }

        # Process all documents concurrently with semaphore limiting
        tasks = [process_with_limit(fp) for fp in file_paths]
        results = await asyncio.gather(*tasks, return_exceptions=True)

        # Convert exceptions to error results
        processed_results = []
        for i, result in enumerate(results):
            if isinstance(result, Exception):
                logger.error(f"Task {i} failed with exception: {result}")
                processed_results.append(
                    {
                        "error": str(result),
                        "file_name": Path(file_paths[i]).name,
                        "document_type": "unknown",
                        "success": False,
                    }
                )
            else:
                processed_results.append(result)

        return processed_results
